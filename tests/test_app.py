"""Application shell: command line, exports, plugins, project state and translated menus."""
import numpy as np
import pandas as pd
import pytest

from openstat import i18n
from openstat.core import registry
from openstat.core.export import export_docx, parse_output
from openstat.core.result import Result, wrap_document
from openstat.core.worksheet import Worksheet, normalize_frame
from openstat.graphs import new_figure


@pytest.fixture
def window(qapp):
    from openstat.ui.main_window import MainWindow

    w = MainWindow()
    w.add_worksheet(Worksheet("Data", pd.DataFrame({"A": [1.0, 2.0, 3.0, 4.0], "B": [2.0, 4.0, 6.0, 9.0]})))
    return w


def sample_result() -> Result:
    fig = new_figure(3, 2)
    fig.add_subplot().plot([1, 2, 3], [1, 4, 9])
    res = Result("Example").text("first line\nsecond line").heading("Section")
    res.table(pd.DataFrame({"N": [3, 4], "Mean": [1.5, 2.25]}, index=pd.Index(["x", "y"], name="Variable")))
    return res.pre("  1 | 23").figure(fig)


def test_normalize_frame():
    frame = normalize_frame(pd.DataFrame({1: [1, 2], "flag": [True, False], "mixed": [1.5, "a"],
                                          "numbers": pd.Series([1, None], dtype=object)}))
    assert list(frame.columns) == ["1", "flag", "mixed", "numbers"]
    assert frame["1"].dtype == float and frame["flag"].tolist() == [1.0, 0.0]
    assert frame["mixed"].tolist() == ["1.5", "a"]
    assert frame["numbers"].dtype == float and np.isnan(frame["numbers"][1])


def test_console_edits_the_worksheet_and_runs_analyses(window):
    console = window.console
    console.execute("df['C'] = df['A'] * 2")
    assert window.current_worksheet().df["C"].tolist() == [2.0, 4.0, 6.0, 8.0]
    assert console.execute("print(int(df['B'].sum()))") == "21"
    console.execute("run('basic.describe', columns=['A', 'B'])")
    assert len(window.project.outputs) == 1 and window.navigator.count() == 1
    assert "NameError" in console.execute("undefined_name")
    assert "basic.describe" in console.execute("analyses('descr')")


def test_console_drops_properties_of_removed_columns(window):
    ws = window.current_worksheet()
    ws.descriptions["A"], ws.value_orders["A"], ws.formats["A"] = "d", "occurrence", [{"rule": "x"}]
    ws.descriptions["B"] = "kept"
    window.console.execute("df = df.drop(columns=['A'])")
    assert (ws.descriptions, ws.value_orders, ws.formats) == ({"B": "kept"}, {}, {})


def test_formulas_dates_and_worksheet_parameters(window):
    from PySide6.QtCore import Qt

    from openstat.ui.dialog_builder import AnalysisDialog

    console, ws, model = window.console, window.current_worksheet(), window.current_view().model_
    console.execute("run('calc.formula', store='Total', expression='A + B')")
    assert ws.df["Total"].tolist() == [3.0, 6.0, 9.0, 13.0]
    model.set_cells([(0, 0, "10")])                       # editing the grid recalculates
    console.execute("df['B'] = 0")                        # and so does the console
    assert ws.df["Total"].tolist() == [10.0, 2.0, 3.0, 4.0]
    assert model.headerData(2, Qt.Horizontal).startswith("C3 =")
    assert model.headerData(2, Qt.Horizontal, Qt.ToolTipRole) == "= A + B"
    assert "own column" in console.execute("run('calc.formula', store='Total', expression='Total + 1')")

    model.set_cells([(0, 3, "2024-05-01"), (1, 3, "2024-05-02")])
    assert model.headerData(3, Qt.Horizontal).startswith("C4-D") and model.data(model.index(1, 3)) == "2024-05-02"

    first = window.project.worksheets[0]
    dialog = AnalysisDialog(registry.get("data.merge"), ws, worksheets=window.project.worksheets)
    assert dialog.values()["other"] is first and dialog.values()["by"] == []
    dialog.accept()                                       # the matching columns are optional
    assert dialog.result() == AnalysisDialog.Accepted
    console.execute("run('data.merge', other='Data')")
    assert window.current_worksheet().name == "Data + Data" and window.current_worksheet().ncols == 2 * ws.ncols


def test_console_multiline_and_history(window):
    console = window.console
    assert console.execute("for i in range(2):") == "" and console.pending
    console.execute("    print(i)")
    assert console.execute("") == "0\n1" and not console.pending
    console.line.remember("a = 1")
    console.line.remember("a = 1")
    assert console.line.history == ["a = 1"]


def test_parse_output_blocks():
    kinds = [kind for kind, _ in parse_output(sample_result().to_html())]
    assert kinds == ["h2", "p", "h3", "table", "pre", "image"]
    blocks = dict(parse_output(sample_result().to_html()))
    assert blocks["table"][0] == ["Variable", "N", "Mean"] and blocks["table"][1][0] == "x"
    assert blocks["p"] == "first line\nsecond line" and blocks["image"][:4] == b"\x89PNG"


def test_exports_write_real_files(window, tmp_path):
    from docx import Document

    window.add_result(sample_result(), window.current_worksheet())
    for kind in ("html", "docx", "pdf"):
        window.export_output(kind, str(tmp_path / f"out.{kind}"))
    assert "<table" in (tmp_path / "out.html").read_text(encoding="utf-8")
    assert (tmp_path / "out.pdf").read_bytes()[:5] == b"%PDF-"
    document = Document(str(tmp_path / "out.docx"))
    assert document.tables[0].cell(1, 2).text == "1.5000"
    assert len(document.inline_shapes) == 1
    assert [p.text for p in document.paragraphs if p.style.name.startswith("Heading")] == ["Example", "Section"]
    export_docx(wrap_document(""), tmp_path / "empty.docx")


def test_plugins_register_analyses_and_report_failures(tmp_path):
    (tmp_path / "good.py").write_text(
        "from openstat.core.registry import analysis, Param\n"
        "from openstat.core.result import Result\n"
        "@analysis('plugin.total', menu=('Plugins',), title='Total', params=[Param('column', 'Column', 'column')])\n"
        "def total(ws, column):\n"
        "    return Result('Total').text(str(ws.df[column].sum()))\n", encoding="utf-8")
    (tmp_path / "broken.py").write_text("raise RuntimeError('boom')\n", encoding="utf-8")
    try:
        errors = registry.load_plugins(tmp_path)
        assert errors == ["broken.py: RuntimeError: boom"]
        ws = Worksheet(df=pd.DataFrame({"A": [1.0, 2.0]}))
        assert registry.get("plugin.total").func(ws, "A").blocks[0] == ("text", "3.0")
    finally:
        registry._REGISTRY.pop("plugin.total", None)
    assert registry.load_plugins(tmp_path / "missing") == []


def test_dirty_state_and_project_roundtrip(window, tmp_path):
    from openstat.core.project import Project

    window.set_project(Project())
    assert not window.dirty and "*" not in window.windowTitle()
    window.current_view().model_.set_cells([(0, 0, "5"), (1, 0, "7")])
    assert window.dirty and "*" in window.windowTitle()
    window.add_result(sample_result(), window.current_worksheet())
    window._save(tmp_path / "p.ostat")
    assert not window.dirty and window.windowTitle().startswith("p.ostat")
    window.set_project(Project.load(tmp_path / "p.ostat"))
    assert window.current_worksheet().df["C1"].tolist() == [5.0, 7.0]
    assert window.navigator.count() == 1 and "Example" in window.output.toPlainText()


def test_menus_are_translated(qapp):
    from openstat.ui.main_window import MainWindow

    i18n.set_language("pt_BR")
    try:
        window = MainWindow()
        texts = set()

        def collect(menu):
            for action in menu.actions():
                texts.add(action.text())
                if action.menu():
                    collect(action.menu())

        for action in window.menuBar().actions():
            texts.add(action.text())
            collect(action.menu())
        assert {"Arquivo", "Estatística", "Ferramentas", "Estatística descritiva...", "Cartas de controle",
                "Carta Xbarra-R...", "Linha de comando", "Exportar saída"} <= texts
    finally:
        i18n.set_language("en")

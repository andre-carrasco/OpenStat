"""Exporting the output to Word; HTML export lives in result.wrap_document and PDF in the UI."""
from __future__ import annotations

import base64
import io
from html.parser import HTMLParser


class _OutputParser(HTMLParser):
    """Turns the limited HTML produced by Result.to_html into a flat list of blocks."""

    def __init__(self):
        super().__init__(convert_charrefs=True)
        self.blocks: list[tuple] = []
        self._text: list[str] = []
        self._kind: str | None = None
        self._rows: list[list[str]] | None = None
        self._cell: list[str] | None = None

    def handle_starttag(self, tag, attrs):
        attrs = dict(attrs)
        if tag in ("h2", "h3", "p", "pre"):
            self._kind, self._text = tag, []
        elif tag == "br" and self._kind:
            self._text.append("\n")
        elif tag == "table":
            self._rows = []
        elif tag == "tr" and self._rows is not None:
            self._rows.append([])
        elif tag in ("td", "th") and self._rows is not None:
            self._cell = []
        elif tag == "img" and attrs.get("src", "").startswith("data:image"):
            self.blocks.append(("image", base64.b64decode(attrs["src"].split(",", 1)[1])))
        elif tag == "hr":
            self.blocks.append(("rule", None))

    def handle_endtag(self, tag):
        if tag in ("h2", "h3", "p", "pre") and self._kind == tag:
            text = "".join(self._text)
            if text.strip():
                self.blocks.append((tag, text if tag == "pre" else text.strip()))
            self._kind = None
        elif tag in ("td", "th") and self._cell is not None and self._rows:
            self._rows[-1].append("".join(self._cell).strip())
            self._cell = None
        elif tag == "table" and self._rows is not None:
            self.blocks.append(("table", [r for r in self._rows if r]))
            self._rows = None

    def handle_data(self, data):
        if self._cell is not None:
            self._cell.append(data)
        elif self._kind:
            self._text.append(data)


def parse_output(html: str) -> list[tuple]:
    parser = _OutputParser()
    parser.feed(html)
    return parser.blocks


def export_docx(html: str, path) -> None:
    """Write the output history as a Word document."""
    from docx import Document
    from docx.shared import Inches, Pt

    document = Document()
    for kind, payload in parse_output(html):
        if kind == "h2":
            document.add_heading(payload, level=1)
        elif kind == "h3":
            document.add_heading(payload, level=2)
        elif kind == "p":
            document.add_paragraph(payload)
        elif kind == "pre":
            run = document.add_paragraph().add_run(payload)
            run.font.name, run.font.size = "Consolas", Pt(9)
        elif kind == "table" and payload:
            width = max(len(row) for row in payload)
            table = document.add_table(rows=len(payload), cols=width)
            table.style = "Light Grid Accent 1"
            for i, row in enumerate(payload):
                for j, text in enumerate(row):
                    cell = table.cell(i, j)
                    cell.text = text
                    for paragraph in cell.paragraphs:
                        for run in paragraph.runs:
                            run.font.size = Pt(9)
                            run.font.bold = i == 0
            document.add_paragraph()
        elif kind == "image":
            document.add_picture(io.BytesIO(payload), width=Inches(6.0))
    document.save(str(path))
